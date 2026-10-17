"""Render a case report as a Word document (.docx) with python-docx. The
content comes from ``app.reports.outline``, so it matches the PDF."""

import io
from datetime import UTC, datetime
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from app.db.models import AuditEvent, Case, EvidenceFile, User
from app.reports import outline


class _Docx:
    def __init__(self, footer_text: str) -> None:
        self.doc = Document()
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        footer = self.doc.sections[0].footer.paragraphs[0]
        footer.text = footer_text
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = _grey(110)

    def heading(self, text: str, size: float = 13) -> None:
        paragraph = self.doc.add_heading(level=1 if size > 15 else 2)
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.font.color.rgb = _grey(20)

    def para(self, text: str, size: float = 10, bold: bool = False, color: int = 40) -> None:
        run = self.doc.add_paragraph().add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = _grey(color)

    def row(self, label: str, value: Any) -> None:
        paragraph = self.doc.add_paragraph()
        name = paragraph.add_run(f"{label}: ")
        name.bold = True
        name.font.color.rgb = _grey(70)
        paragraph.add_run("" if value is None else str(value))

    def gap(self, mm: float) -> None:
        self.doc.add_paragraph().paragraph_format.space_after = Pt(mm * 2)

    def output(self) -> bytes:
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()


def _grey(level: int) -> RGBColor:
    return RGBColor(level, level, level)


def render(
    case: Case,
    options: dict[str, Any],
    evidence: list[EvidenceFile],
    export: dict[str, Any],
    timeline: list[tuple[AuditEvent, User | None]],
) -> bytes:
    generated = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    doc = _Docx(outline.footer_text(case, generated))
    outline.write(doc, case, options, evidence, export, timeline, generated)
    return doc.output()
