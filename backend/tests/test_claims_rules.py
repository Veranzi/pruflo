"""Claims, grounding, rules and explanations (Steps 9-12).

Documents here are built in the test with the same libraries suppliers use
(a PDF with a text layer, an Excel workbook, a Word file), then go through
the real upload, parsing, extraction, grounding and rule path.
"""

import io
import json
from datetime import date
from decimal import Decimal
from typing import Any

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models import Case, ExtractedClaim, Finding, ProviderSettings
from app.enums import CaseStatus, GroundingStatus, Role, RuleTemplate
from app.pipeline.claims import extractor_for, gemini
from app.pipeline.claims.schema import (
    DocumentType,
    RawClaim,
    parse_date,
    parse_money,
    parse_quantity,
    validate,
)
from app.pipeline.claims.structured import StructuredExtractor
from app.pipeline.document import BBox, Block, Page
from app.pipeline.errors import PermanentError, TransientError
from app.pipeline.grounding import ground
from app.pipeline.rules.engine import Doc, Fact, evaluate
from tests import samples
from tests.factories import Factory, Person
from tests.test_evidence import upload

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _page(lines: list[str], index: int = 1, boxes: bool = False) -> Page:
    blocks = [
        Block(text=t, bbox=BBox(0.1, 0.05 * n, 0.9, 0.05 * n + 0.04) if boxes else None)
        for n, t in enumerate(lines, start=1)
    ]
    page = Page(index=index, blocks=blocks)
    page.assign_spans()
    return page


# Parsing values (code, never the model) --------------------------------------------


def test_values_are_parsed_by_code() -> None:
    assert parse_money("SGD 1,250.50") == (Decimal("1250.50"), "SGD")
    assert parse_money("RM 3,000") == (Decimal("3000"), "MYR")
    assert parse_money("Rp 1.500.000") == (Decimal("1500000"), "IDR")
    assert parse_money("$ 99.90") == (Decimal("99.90"), None)  # "$" alone is ambiguous
    assert parse_quantity("1,200 pcs") == (Decimal("1200"), "pcs")
    assert parse_quantity("80 Units") == (Decimal("80"), "units")
    assert parse_quantity("about 80") == (None, None)
    assert parse_date("03/10/2026") == date(2026, 10, 3)  # day first in SEA
    assert parse_date("03/10/2026", day_first=False) == date(2026, 3, 10)
    assert parse_date("3 Oct 2026") == date(2026, 10, 3)
    assert parse_date("October 3, 2026") == date(2026, 10, 3)
    assert parse_date("31/02/2026") is None  # impossible dates aren't guessed


def test_validation_rejects_anything_outside_the_schema() -> None:
    assert validate(RawClaim("drop_table", "1", "x 1")) is None
    assert validate(RawClaim("ordered_quantity", "100", "Qty 100")) is None  # no item
    assert validate(RawClaim("document_type", "memo", "Memo")) is None
    claim = validate(RawClaim("ordered_quantity", "about 100", "Bolts about 100", item="Bolts"))
    assert claim is not None and claim.status == "unreadable" and claim.normalized_value is None
    long = validate(RawClaim("supplier_name", "x" * 1000, "Supplier: " + "x" * 1000))
    assert long is not None and len(long.raw_value) == 200 and len(long.quote) == 500


# Structured extraction ---------------------------------------------------------------


PO_LINES = [
    "ACME BUILDING SUPPLIES PTE LTD",
    "PURCHASE ORDER",
    "PO No: PO-2026-0042 Date: 01/10/2026",
    "Supplier: Widget Hardware Pte Ltd",
    "Delivery Date: 05/10/2026",
    "Item Description Qty Unit Unit Price Amount",
    "1 Steel bolts M8 100 pcs 0.50 50.00",
    "2 Washers M8 200 pcs 0.10 20.00",
    "Subtotal 70.00",
    "GST 9% 6.30",
    "Grand Total: SGD 76.30",
]


def test_structured_reads_a_purchase_order() -> None:
    claims = {(c.field, c.item): c for c in StructuredExtractor().extract([_page(PO_LINES)])}
    assert claims[("document_type", None)].raw_value == "purchase_order"
    assert claims[("document_number", None)].raw_value == "PO-2026-0042"
    assert claims[("promised_delivery_date", None)].raw_value == "05/10/2026"
    assert claims[("total_amount", None)].raw_value == "SGD 76.30"  # not the subtotal
    bolts = claims[("ordered_quantity", "Steel bolts M8")]
    assert bolts.raw_value == "100" and bolts.unit == "pcs"
    assert bolts.quote == "1 Steel bolts M8 100 pcs 0.50 50.00"
    assert claims[("ordered_quantity", "Washers M8")].raw_value == "200"


def test_structured_skips_references_to_other_documents_in_the_title() -> None:
    lines = ["Purchase Order No: PO-9", "DELIVERY ORDER", "DO No: DO-1"]
    kinds = [
        c.raw_value
        for c in StructuredExtractor().extract([_page(lines)])
        if c.field == "document_type"
    ]
    assert kinds == ["delivery_note"]


def test_structured_claims_nothing_from_prose() -> None:
    lines = ["Meeting notes", "We discussed the item quantity for next quarter."]
    assert StructuredExtractor().extract([_page(lines)]) == []


# Grounding (golden cases) --------------------------------------------------------------


def test_grounding_exact_quote_gets_span_and_box() -> None:
    page = _page(["PURCHASE ORDER", "Qty ordered: 100", "Total: SGD 50"], boxes=True)
    result = ground("Qty ordered: 100", "100", 1, [page])
    assert result.status == GroundingStatus.GROUNDED and result.page == 1
    assert page.text[result.span_start : result.span_end] == "Qty ordered: 100"
    assert result.bbox == {"x0": 0.1, "y0": 0.1, "x1": 0.9, "y1": 0.14}


def test_grounding_ignores_only_whitespace_and_line_wraps() -> None:
    page = _page(["Goods were deliv-", "ered on 08/10/2026  in full"])
    assert (
        ground("Goods were delivered on 08/10/2026 in full", "08/10/2026", 1, [page]).status
        == GroundingStatus.GROUNDED
    )
    # Any other difference (case, a changed digit) isn't a match.
    assert (
        ground("goods were delivered on 08/10/2026", "08/10/2026", 1, [page]).status
        == GroundingStatus.UNGROUNDED
    )


def test_grounding_missing_quote_is_ungrounded() -> None:
    page = _page(["Qty ordered: 100"])
    assert ground("Qty ordered: 120", "120", 1, [page]).status == GroundingStatus.UNGROUNDED
    # The quote must contain the value it supports.
    assert ground("Qty ordered: 100", "120", 1, [page]).status == GroundingStatus.UNGROUNDED


def test_grounding_repeated_quote_is_ambiguous() -> None:
    page = _page(["Washers 200 pcs", "Bolts 100 pcs", "Washers 200 pcs"])
    assert ground("Washers 200 pcs", "200", 1, [page]).status == GroundingStatus.AMBIGUOUS


def test_grounding_bare_number_is_never_trusted() -> None:
    page = _page(["Delivered", "80"])
    assert ground("80", "80", 1, [page]).status == GroundingStatus.AMBIGUOUS


def test_grounding_uses_the_page_hint_then_the_whole_document() -> None:
    pages = [_page(["Total: SGD 50"], 1), _page(["Total: SGD 50"], 2)]
    assert ground("Total: SGD 50", "SGD 50", 2, pages).page == 2
    assert ground("Total: SGD 50", "SGD 50", 3, pages).status == GroundingStatus.AMBIGUOUS
    only_two = [_page(["Header"], 1), _page(["Total: SGD 50"], 2)]
    found = ground("Total: SGD 50", "SGD 50", 1, only_two)
    assert found.status == GroundingStatus.GROUNDED and found.page == 2


# Rule engine (no model involved) -----------------------------------------------------------


def _fact(
    n: int,
    field: str,
    value: str,
    unit: str | None = None,
    item: str | None = None,
    file: str = "po.pdf",
    evidence: str = "e1",
) -> Fact:
    return Fact(f"c{n}", evidence, file, field, value, unit, item, item or "")


ALL_RULES = {
    RuleTemplate.QUANTITY_SHORTFALL: {"tolerance_units": "0"},
    RuleTemplate.AMOUNT_MISMATCH: {"tolerance_percent": "0"},
    RuleTemplate.DELIVERY_DATE_WINDOW: {"grace_days": "0"},
    RuleTemplate.MISSING_DOCUMENT: {},
}


def _docs(delivered: str = "80", dn_unit: str | None = "pcs") -> list[Doc]:
    po = Doc(
        "e1",
        "po.pdf",
        DocumentType.PURCHASE_ORDER,
        [
            _fact(1, "document_type", "purchase_order"),
            _fact(2, "document_number", "PO-1"),
            _fact(3, "ordered_quantity", "100", "pcs", "widget"),
            _fact(4, "total_amount", "1000", "SGD"),
            _fact(5, "promised_delivery_date", "2026-10-05"),
        ],
    )
    dn = Doc(
        "e2",
        "dn.xlsx",
        DocumentType.DELIVERY_NOTE,
        [
            _fact(6, "document_type", "delivery_note", file="dn.xlsx", evidence="e2"),
            _fact(7, "order_reference", "po-1", file="dn.xlsx", evidence="e2"),
            _fact(8, "delivered_quantity", delivered, dn_unit, "widget", "dn.xlsx", "e2"),
            _fact(9, "delivery_date", "2026-10-08", file="dn.xlsx", evidence="e2"),
        ],
    )
    inv = Doc(
        "e3",
        "inv.docx",
        DocumentType.INVOICE,
        [
            _fact(10, "document_type", "invoice", file="inv.docx", evidence="e3"),
            _fact(11, "total_amount", "1100", "SGD", file="inv.docx", evidence="e3"),
        ],
    )
    return [po, dn, inv]


def test_shortfall_of_20_is_computed_in_code() -> None:
    found = {p.finding_type: p for p in evaluate(_docs(), ALL_RULES)}
    shortfall = found["quantity_shortfall"]
    assert shortfall.expected["value"] == "100" and shortfall.observed["value"] == "80"
    assert shortfall.difference == {"value": "20", "unit": "pcs"}
    assert shortfall.severity == "high"
    assert {"c3", "c8"} <= set(shortfall.claim_ids)
    assert (
        "lists 100 pcs of widget" in shortfall.explanation
        and "80 pcs, 20 pcs short" in shortfall.explanation
    )
    assert found["amount_mismatch"].difference == {"value": "100", "unit": "SGD", "percent": "10"}
    assert found["late_delivery"].difference == {"value": "3", "unit": "days"}
    assert "missing_document" not in found


def test_tolerances_and_grace_come_from_the_rule_parameters() -> None:
    params = {
        RuleTemplate.QUANTITY_SHORTFALL: {"tolerance_units": "20"},
        RuleTemplate.AMOUNT_MISMATCH: {"tolerance_percent": "10"},
        RuleTemplate.DELIVERY_DATE_WINDOW: {"grace_days": "3"},
    }
    assert evaluate(_docs(), params) == []
    assert evaluate(_docs(), {}) == []  # every rule switched off


def test_different_units_are_never_compared() -> None:
    found = evaluate(_docs(dn_unit="boxes"), {RuleTemplate.QUANTITY_SHORTFALL: {}})
    assert [p.finding_type for p in found] == ["data_quality"]
    assert "never converted" in found[0].explanation


def test_missing_documents_are_reported_as_absence_of_evidence() -> None:
    po = _docs()[0]
    found = evaluate([po], ALL_RULES)
    missing = sorted(p.expected["value"] for p in found if p.finding_type == "missing_document")
    assert missing == ["delivery note", "invoice"]
    assert all("absence of evidence" in p.explanation for p in found)
    # A delivery note whose order isn't in the case.
    orphan = _docs()[1]
    found = evaluate([orphan], ALL_RULES)
    assert found[0].expected == {"value": "purchase order", "reference": "po-1"}


def test_unreferenced_document_with_several_orders_isnt_guessed() -> None:
    po1, dn, _ = _docs()
    po2 = Doc(
        "e9",
        "po2.pdf",
        DocumentType.PURCHASE_ORDER,
        [_fact(20, "document_number", "PO-2", evidence="e9")],
    )
    dn.facts = [f for f in dn.facts if f.field != "order_reference"]
    kinds = [p.finding_type for p in evaluate([po1, po2, dn], ALL_RULES)]
    assert "data_quality" in kinds and "quantity_shortfall" not in kinds


def test_same_inputs_same_findings() -> None:
    first = [p.dedupe_key("rule_q", 1) for p in evaluate(_docs(), ALL_RULES)]
    second = [p.dedupe_key("rule_q", 1) for p in evaluate(_docs(), ALL_RULES)]
    assert first == second and len(set(first)) == len(first)
    assert evaluate(_docs(), ALL_RULES)[0].dedupe_key("rule_q", 2) not in first


# End to end through the API -------------------------------------------------------------


def _po_pdf() -> bytes:
    return samples.text_pdf([PO_LINES])


def _dn_xlsx(
    bolts: int = 80, washers: int = 200, number: str = "DO-881", day: str = "08/10/2026"
) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Delivery"
    sheet.append(["DELIVERY ORDER"])
    sheet.append(["DO No", number])
    sheet.append(["PO Ref", "PO-2026-0042"])
    sheet.append(["Date", day])
    sheet.append([])
    sheet.append(["Description", "Qty Delivered", "UOM"])
    if bolts:
        sheet.append(["Steel bolts M8", bolts, "pcs"])
    if washers:
        sheet.append(["Washers M8", washers, "pcs"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _invoice_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("TAX INVOICE")
    document.add_paragraph("Invoice No: INV-7781")
    document.add_paragraph("Your Order: PO-2026-0042")
    table = document.add_table(rows=3, cols=3)
    for r, row in enumerate(
        [
            ["Description", "Qty", "Amount"],
            ["Steel bolts M8", "100", "50.00"],
            ["Washers M8", "200", "20.00"],
        ]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph("Total Amount Due: SGD 83.70")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _findings(client: TestClient, person: Person, case_id: str) -> list[dict[str, Any]]:
    res = client.get(f"/v1/cases/{case_id}/findings", headers=person.headers)
    assert res.status_code == 200, res.text
    return list(res.json()["items"])


def test_order_delivery_and_invoice_produce_grounded_findings(
    client: TestClient, make: Factory, db: Session
) -> None:
    alice = make.person(Role.INVESTIGATOR)
    reviewer = make.person(Role.REVIEWER, alice.tenant)
    case = make.case(alice.tenant, alice, reviewer)
    assert upload(client, alice, case.id, "po.pdf", _po_pdf()).status_code == 200
    assert upload(client, alice, case.id, "delivery.xlsx", _dn_xlsx(), XLSX).status_code == 200
    assert upload(client, alice, case.id, "invoice.docx", _invoice_docx(), DOCX).status_code == 200

    claims = client.get(f"/v1/cases/{case.id}/claims", headers=alice.headers).json()["items"]
    po_bolts = next(
        c
        for c in claims
        if c["field"] == "ordered_quantity" and c["file_name"] == "po.pdf" and c["value"] == "100"
    )
    assert po_bolts["grounding"] == "grounded" and po_bolts["page"] == 1
    assert po_bolts["bbox"] is not None  # a PDF line has a box to highlight
    assert po_bolts["quote"] == "1 Steel bolts M8 100 pcs 0.50 50.00"
    assert po_bolts["model"] == "structured-1"

    found = {f["type"]: f for f in _findings(client, alice, case.id)}
    shortfall = found["quantity_shortfall"]
    assert shortfall["expected"] == {"value": "100", "unit": "pcs", "item": "Steel bolts M8"}
    assert shortfall["observed"]["value"] == "80" and shortfall["difference"]["value"] == "20"
    assert shortfall["review"]["status"] == "pending"
    assert shortfall["explanation_source"] == "template"
    assert {c["file_name"] for c in shortfall["evidence"]} >= {"po.pdf", "delivery.xlsx"}
    amount = found["amount_mismatch"]
    assert amount["expected"]["value"] == "76.3" and amount["observed"]["value"] == "83.7"
    assert found["late_delivery"]["difference"] == {"value": "3", "unit": "days"}
    assert "missing_document" not in found
    db.expire_all()
    stored = db.get(Case, case.id)
    assert stored is not None and stored.status == CaseStatus.IN_REVIEW

    # People decide; when nothing is pending the case is completed.
    for finding in found.values():
        res = client.post(
            f"/v1/cases/{case.id}/findings/{finding['id']}/decision",
            headers=reviewer.headers,
            json={"decision": "approve"},
        )
        assert res.status_code == 200, res.text
    db.expire_all()
    stored = db.get(Case, case.id)
    assert stored is not None and stored.status == CaseStatus.COMPLETED


def test_a_later_delivery_clears_an_unreviewed_shortfall(
    client: TestClient, make: Factory, db: Session
) -> None:
    alice = make.person(Role.INVESTIGATOR)
    case = make.case(alice.tenant, alice)
    upload(client, alice, case.id, "po.pdf", _po_pdf())
    upload(client, alice, case.id, "delivery.xlsx", _dn_xlsx(), XLSX)
    assert "quantity_shortfall" in {f["type"] for f in _findings(client, alice, case.id)}
    # The remaining 20 bolts arrive on a second delivery order, on time.
    second = _dn_xlsx(bolts=20, washers=0, number="DO-882", day="04/10/2026")
    upload(client, alice, case.id, "delivery-2.xlsx", second, XLSX)
    types = {f["type"] for f in _findings(client, alice, case.id)}
    assert "quantity_shortfall" not in types
    assert (
        db.scalar(
            select(Finding).where(
                Finding.case_id == case.id, Finding.finding_type == "quantity_shortfall"
            )
        )
        is None
    )


def test_reviewed_findings_are_kept_and_rules_follow_parameters(
    client: TestClient, make: Factory, db: Session
) -> None:
    admin = make.person(Role.ADMIN)
    rules = client.get("/v1/admin/rules", headers=admin.headers).json()["items"]
    qty = next(r for r in rules if r["template"] == "quantity_shortfall")
    res = client.patch(
        f"/v1/admin/rules/{qty['id']}",
        headers=admin.headers,
        json={"params": {"tolerance_units": "25"}},
    )
    assert res.status_code == 200 and res.json()["version"] == 2
    case = make.case(admin.tenant, admin)
    upload(client, admin, case.id, "po.pdf", _po_pdf())
    upload(client, admin, case.id, "delivery.xlsx", _dn_xlsx(), XLSX)
    # A 20-unit gap is within the 25-unit tolerance.
    assert "quantity_shortfall" not in {f["type"] for f in _findings(client, admin, case.id)}


def test_scanned_or_unreadable_values_stay_out_of_the_rules(
    client: TestClient, make: Factory, db: Session
) -> None:
    alice = make.person(Role.INVESTIGATOR)
    case = make.case(alice.tenant, alice)
    lines = [ln for ln in PO_LINES if "Washers" not in ln] + ["1 Steel bolts M8 100 pcs 0.50 50.00"]
    upload(client, alice, case.id, "po.pdf", samples.text_pdf([lines]))
    claims = db.scalars(select(ExtractedClaim).where(ExtractedClaim.case_id == case.id)).all()
    bolts = [c for c in claims if c.field_name == "ordered_quantity"]
    # The same row appears twice, so its quote is ambiguous and not used by rules.
    assert bolts and all(c.grounding == GroundingStatus.AMBIGUOUS for c in bolts)


# Gemini adapter (against a stand-in for Vertex AI) ------------------------------------------


@pytest.fixture
def vertex(monkeypatch: pytest.MonkeyPatch) -> list[httpx.Request]:
    calls: list[httpx.Request] = []
    reply: dict[str, Any] = {
        "claims": [
            {
                "field_name": "document_type",
                "raw_value": "purchase_order",
                "source_quote": "PURCHASE ORDER",
                "page_hint": 1,
                "extraction_status": "found",
            },
            {
                "field_name": "ordered_quantity",
                "raw_value": "100",
                "item": "Steel bolts M8",
                "unit": "pcs",
                "source_quote": "1 Steel bolts M8 100 pcs 0.50 50.00",
                "page_hint": 1,
                "extraction_status": "found",
            },
            {
                "field_name": "total_amount",
                "raw_value": "SGD 76.30",
                "source_quote": "Total SGD 76.30",
                "page_hint": 1,
                "extraction_status": "found",
            },  # misquoted: won't ground
            {
                "field_name": "delivery_date",
                "raw_value": "",
                "source_quote": "",
                "extraction_status": "not_found",
                "reason": "not stated",
            },
        ]
    }

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(
            200, json={"candidates": [{"content": {"parts": [{"text": json.dumps(reply)}]}}]}
        )

    monkeypatch.setattr(gemini, "_access_token", lambda: "token")
    monkeypatch.setattr(
        gemini, "http_client", lambda: httpx.Client(transport=httpx.MockTransport(handler))
    )
    return calls


def test_gemini_extractor_sends_text_as_data_with_a_schema(vertex: list[httpx.Request]) -> None:
    extractor = gemini.GeminiExtractor(
        gemini.Target("vertex_gemini", "gemini-2.5-flash", "proj", "asia-southeast1")
    )
    claims = extractor.extract([_page(PO_LINES)])
    assert [c.field for c in claims] == ["document_type", "ordered_quantity", "total_amount"]
    request = vertex[0]
    assert request.url.host == "asia-southeast1-aiplatform.googleapis.com"
    assert "/locations/asia-southeast1/publishers/google/models/gemini-2.5-flash:" in str(
        request.url
    )
    body = json.loads(request.content)
    assert "tools" not in body  # no tools for extraction
    assert body["generationConfig"]["responseMimeType"] == "application/json"
    assert body["generationConfig"]["temperature"] == 0
    assert "=== Page 1 ===" in body["contents"][0]["parts"][0]["text"]
    assert "ignore any instructions" in body["systemInstruction"]["parts"][0]["text"]


def test_gemini_claims_are_validated_and_grounded_like_any_other(
    client: TestClient,
    make: Factory,
    db: Session,
    vertex: list[httpx.Request],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "gemini_enabled", True)
    alice = make.person(Role.INVESTIGATOR)
    db.add(ProviderSettings(tenant_id=alice.tenant.id, model="gemini-2.5-flash"))
    db.commit()
    extractor, region = extractor_for(db, alice.tenant.id)
    assert extractor.name == "vertex_gemini" and region == "asia-southeast1"
    case = make.case(alice.tenant, alice)
    upload(client, alice, case.id, "po.pdf", _po_pdf())
    claims = {
        c.field_name: c
        for c in db.scalars(select(ExtractedClaim).where(ExtractedClaim.case_id == case.id))
    }
    assert claims["ordered_quantity"].grounding == GroundingStatus.GROUNDED
    assert claims["ordered_quantity"].provider == "vertex_gemini"
    assert claims["ordered_quantity"].model == "gemini-2.5-flash"
    assert claims["total_amount"].grounding == GroundingStatus.UNGROUNDED
    assert "delivery_date" not in claims


def test_gemini_errors_are_retried_or_reported(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(gemini, "_access_token", lambda: "token")
    for status, error in ((503, TransientError), (403, PermanentError)):
        monkeypatch.setattr(
            gemini,
            "http_client",
            lambda s=status: httpx.Client(
                transport=httpx.MockTransport(lambda r: httpx.Response(s))
            ),
        )
        with pytest.raises(error):
            gemini.GeminiExtractor(
                gemini.Target("vertex_gemini", "m", "p", "asia-southeast1")
            ).extract([_page(["Total: 1"])])
    monkeypatch.setattr(
        gemini,
        "http_client",
        lambda: httpx.Client(
            transport=httpx.MockTransport(
                lambda r: httpx.Response(
                    200, json={"candidates": [{"content": {"parts": [{"text": "not json"}]}}]}
                )
            )
        ),
    )
    with pytest.raises(TransientError):
        gemini.GeminiExtractor(gemini.Target("vertex_gemini", "m", "p", "asia-southeast1")).extract(
            [_page(["Total: 1"])]
        )


def test_structured_extractor_is_used_without_gemini(make: Factory, db: Session) -> None:
    person = make.person(Role.INVESTIGATOR)
    extractor, region = extractor_for(db, person.tenant.id)
    assert extractor.name == "pruflo_structured" and region == get_settings().region


# Case Q&A (endpoint 20) --------------------------------------------------------------------


def test_assistant_is_unavailable_without_gemini(client: TestClient, make: Factory) -> None:
    alice = make.person(Role.INVESTIGATOR)
    case = make.case(alice.tenant, alice)
    res = client.post(
        f"/v1/cases/{case.id}/qa",
        headers=alice.headers,
        json={"question": "Why is this case in review?"},
    )
    assert res.status_code == 503 and res.json()["error"]["code"] == "assistant_unavailable"


def test_assistant_answers_with_checked_citations(
    client: TestClient, make: Factory, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.services import assistant

    alice = make.person(Role.VIEWER)
    stranger = make.person(Role.ADMIN)
    case = make.case(alice.tenant, alice)
    finding = make.finding(case)
    other_case = make.case(stranger.tenant, stranger)
    foreign = make.finding(other_case)
    db.add(ProviderSettings(tenant_id=alice.tenant.id, model="gemini-2.5-flash"))
    db.commit()
    monkeypatch.setattr(get_settings(), "gemini_enabled", True)
    turns: list[list[dict[str, Any]]] = []

    def fake_generate(
        target: gemini.Target,
        system: str,
        contents: list[dict[str, Any]],
        config: dict[str, Any],
        tools: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        turns.append(contents)
        assert tools is not None and target.region == "asia-southeast1"
        if len(turns) == 1:
            return {
                "candidates": [
                    {"content": {"parts": [{"functionCall": {"name": "get_findings", "args": {}}}]}}
                ]
            }
        # The tool result came back as data, scoped to this case only.
        result = contents[-1]["parts"][0]["functionResponse"]["response"]["content"]
        assert [f["id"] for f in result] == [finding.id]
        return {
            "candidates": [
                {
                    "content": {
                        "parts": [
                            {
                                "text": f"The delivery was 20 units short [{finding.id}]; "
                                f"see also [{foreign.id}]."
                            }
                        ]
                    }
                }
            ]
        }

    monkeypatch.setattr(assistant.gemini, "generate", fake_generate)
    res = client.post(
        f"/v1/cases/{case.id}/qa", headers=alice.headers, json={"question": "What was short?"}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["citations"] == [{"type": "finding", "id": finding.id}]
    assert foreign.id not in body["answer"] and body["source"] == "machine_generated"
    # Not a member: the case doesn't exist for them.
    outsider = make.person(Role.INVESTIGATOR, alice.tenant)
    res = client.post(
        f"/v1/cases/{case.id}/qa", headers=outsider.headers, json={"question": "What was short?"}
    )
    assert res.status_code == 404


def test_assistant_answer_without_citations_is_withheld(
    client: TestClient, make: Factory, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.services import assistant

    alice = make.person(Role.INVESTIGATOR)
    case = make.case(alice.tenant, alice)
    db.add(ProviderSettings(tenant_id=alice.tenant.id, model="gemini-2.5-flash"))
    db.commit()
    monkeypatch.setattr(get_settings(), "gemini_enabled", True)
    monkeypatch.setattr(
        assistant.gemini,
        "generate",
        lambda *a, **k: {
            "candidates": [{"content": {"parts": [{"text": "Probably the supplier's fault."}]}}]
        },
    )
    res = client.post(
        f"/v1/cases/{case.id}/qa", headers=alice.headers, json={"question": "Whose fault is it?"}
    )
    assert res.json() == {
        "answer": assistant.NO_ANSWER,
        "citations": [],
        "source": "machine_generated",
    }


def test_gemini_api_key_target(
    make: Factory, db: Session, vertex: list[httpx.Request], monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "gemini_enabled", True)
    person = make.person(Role.INVESTIGATOR)
    db.add(ProviderSettings(tenant_id=person.tenant.id, model="gemini-2.5-flash", ai="gemini_api"))
    db.commit()
    # No key configured: Gemini isn't used, the structured reader is.
    assert extractor_for(db, person.tenant.id)[0].name == "pruflo_structured"
    monkeypatch.setattr(settings, "gemini_api_key", "test-key")
    extractor, where = extractor_for(db, person.tenant.id)
    assert extractor.name == "gemini_api" and where == "gemini-api"
    extractor.extract([_page(PO_LINES)])
    request = vertex[-1]
    assert request.url.host == "generativelanguage.googleapis.com"
    assert request.url.path == "/v1beta/models/gemini-2.5-flash:generateContent"
    assert request.headers["x-goog-api-key"] == "test-key"
    assert "authorization" not in request.headers
