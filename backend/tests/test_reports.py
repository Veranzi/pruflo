"""Reports (17-19) and the findings export (50)."""

import csv
import io
from typing import Any
from urllib.parse import urlsplit

from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.enums import Role
from app.worker.queue import Queue
from tests.factories import Factory, Person, key_headers


def _pdf_text(client: TestClient, url: str) -> str:
    parts = urlsplit(url)
    body = client.get(f"{parts.path}?{parts.query}").content
    assert body.startswith(b"%PDF")
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(body)).pages)


def _decided_case(client: TestClient, make: Factory) -> tuple[Person, Any]:
    reviewer = make.person(Role.REVIEWER, name="Rita Reviewer")
    case = make.case(reviewer.tenant, reviewer, title="Supplier shortfall – Nairobi ₹ check")
    case.status = "in_review"
    make.db.commit()
    finding = make.finding(case)
    res = client.post(
        f"/v1/cases/{case.id}/findings/{finding.id}/decision",
        headers=reviewer.headers,
        json={"decision": "reject", "note": "Wrong delivery note attached."},
    )
    assert res.status_code == 200
    return reviewer, case


def test_generate_list_and_download_detailed_report(client: TestClient, make: Factory) -> None:
    reviewer, case = _decided_case(client, make)
    res = client.post(
        f"/v1/cases/{case.id}/reports", headers=reviewer.headers, json={"type": "detailed"}
    )
    assert res.status_code == 202
    assert res.json()["status"] == "queued" and res.json()["type"] == "detailed"

    items = client.get(f"/v1/cases/{case.id}/reports", headers=reviewer.headers).json()["items"]
    assert items[0]["status"] == "ready"  # the inline worker ran after the response
    assert items[0]["requested_by"]["name"] == "Rita Reviewer"

    link = client.get(
        f"/v1/cases/{case.id}/reports/{items[0]['id']}/download-url", headers=reviewer.headers
    ).json()["url"]
    text = _pdf_text(client, link)
    for expected in (
        case.reference,
        "Supplier shortfall – Nairobi ₹ check",
        "Quantity shortfall",
        "100 units",
        "80 units",
        "20 units",
        "reject by Rita Reviewer",
        "Wrong delivery note attached.",
        "Qty ordered: 100",
        "Timeline",
        "human interpretation",
        "Recommendations",
        "Findings & Conclusion",
        "Conclusion: 1 finding(s)",
    ):
        assert expected in text, expected

    activity = client.get("/v1/me/activity", headers=reviewer.headers).json()["items"]
    assert "report_generated" in {a["kind"] for a in activity}


def test_word_report_has_the_same_content(client: TestClient, make: Factory) -> None:
    from docx import Document

    reviewer, case = _decided_case(client, make)
    res = client.post(
        f"/v1/cases/{case.id}/reports",
        headers=reviewer.headers,
        json={"type": "detailed", "format": "docx"},
    )
    assert res.status_code == 202 and res.json()["format"] == "docx"
    report = client.get(f"/v1/cases/{case.id}/reports", headers=reviewer.headers).json()["items"][0]
    assert report["status"] == "ready"

    link = client.get(
        f"/v1/cases/{case.id}/reports/{report['id']}/download-url", headers=reviewer.headers
    ).json()["url"]
    parts = urlsplit(link)
    file = client.get(f"{parts.path}?{parts.query}")
    assert file.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert file.headers["content-disposition"].startswith("attachment;")
    assert ".docx" in file.headers["content-disposition"]
    document = Document(io.BytesIO(file.content))
    text = "\n".join(p.text for p in document.paragraphs)
    header = document.sections[0].header.paragraphs[0]
    assert "PruFlo" in header.text and case.reference in header.text
    assert header._p.xpath(".//pic:pic")  # the logo
    assert "page" in document.sections[0].footer.paragraphs[0].text
    for expected in (
        case.reference,
        "Quantity shortfall",
        "reject by Rita Reviewer",
        "Timeline",
        "Recommendations",
        "Conclusion: 1 finding(s)",
    ):
        assert expected in text, expected


def test_preview_links_open_in_the_browser(client: TestClient, make: Factory) -> None:
    reviewer, case = _decided_case(client, make)
    client.post(f"/v1/cases/{case.id}/reports", headers=reviewer.headers, json={})
    report = client.get(f"/v1/cases/{case.id}/reports", headers=reviewer.headers).json()["items"][0]
    base = f"/v1/cases/{case.id}/reports/{report['id']}/download-url"

    preview = client.get(f"{base}?disposition=inline", headers=reviewer.headers).json()["url"]
    parts = urlsplit(preview)
    shown = client.get(f"{parts.path}?{parts.query}")
    assert shown.headers["content-disposition"].startswith("inline;")
    assert shown.content.startswith(b"%PDF")

    # A download link's file name is signed: changing it breaks the link.
    download = client.get(base, headers=reviewer.headers).json()["url"]
    parts = urlsplit(download)
    assert "dl=" in parts.query
    tampered = parts.query.replace("dl=", "dl=evil-")
    assert client.get(f"{parts.path}?{tampered}").status_code == 403
    assert client.get(f"{parts.path}?{parts.query}").status_code == 200


def test_executive_report_is_short(client: TestClient, make: Factory) -> None:
    reviewer, case = _decided_case(client, make)
    client.post(
        f"/v1/cases/{case.id}/reports",
        headers=reviewer.headers,
        json={"type": "executive", "sections": ["findings"]},
    )
    report = client.get(f"/v1/cases/{case.id}/reports", headers=reviewer.headers).json()["items"][0]
    link = client.get(
        f"/v1/cases/{case.id}/reports/{report['id']}/download-url", headers=reviewer.headers
    ).json()["url"]
    text = _pdf_text(client, link)
    assert "At a glance" in text and "Timeline" not in text
    assert "No findings are waiting for review." in text  # the only finding was decided


def test_download_before_ready_is_409(client: TestClient, make: Factory, queue: Queue) -> None:
    queue.inline = False
    reviewer, case = _decided_case(client, make)
    report = client.post(f"/v1/cases/{case.id}/reports", headers=reviewer.headers, json={}).json()
    res = client.get(
        f"/v1/cases/{case.id}/reports/{report['id']}/download-url", headers=reviewer.headers
    )
    assert res.status_code == 409 and res.json()["error"]["code"] == "report_not_ready"


def test_bad_report_requests(client: TestClient, make: Factory) -> None:
    reviewer, case = _decided_case(client, make)
    for body in ({"type": "everything"}, {"sections": ["secrets"]}):
        res = client.post(f"/v1/cases/{case.id}/reports", headers=reviewer.headers, json=body)
        assert res.status_code == 400, body


def test_json_export_matches_the_contract(client: TestClient, make: Factory) -> None:
    reviewer, case = _decided_case(client, make)
    res = client.get(f"/v1/cases/{case.id}/export?format=json", headers=reviewer.headers)
    assert res.status_code == 200
    assert res.headers["content-disposition"] == (
        f'attachment; filename="case-{case.reference}-findings.json"'
    )
    data = res.json()
    assert data["schema_version"] == "1.0"
    assert data["case"]["reference"] == case.reference
    finding = data["findings"][0]
    assert finding["difference"] == {"value": "20", "unit": "units"}
    assert finding["review"]["status"] == "reject"
    assert finding["review"]["by"] == "Rita Reviewer"
    source = finding["evidence"][0]["source_url"]
    assert source.startswith(f"/v1/cases/{case.id}/evidence/") and "sig=" not in source


def test_csv_export_and_formula_injection(client: TestClient, make: Factory, db: Session) -> None:
    reviewer, case = _decided_case(client, make)
    reviewer.user.name = '=HYPERLINK("http://evil")'
    db.commit()
    res = client.get(f"/v1/cases/{case.id}/export?format=csv", headers=reviewer.headers)
    assert res.headers["content-type"].startswith("text/csv")
    rows = list(csv.DictReader(io.StringIO(res.text)))
    assert len(rows) == 1
    row = rows[0]
    assert (row["expected"], row["observed"], row["difference"], row["unit"]) == (
        "100",
        "80",
        "20",
        "units",
    )
    assert row["review_status"] == "reject"
    assert row["reviewed_by"].startswith("'=")  # neutralised for spreadsheets
    assert row["evidence"] == "purchase_order.pdf p1; delivery_note.pdf p1"


def test_who_may_export(client: TestClient, make: Factory) -> None:
    reviewer, case = _decided_case(client, make)
    viewer = make.person(Role.VIEWER, reviewer.tenant)
    make.db.add_all([])  # viewer must be a member to prove the role (not membership) blocks it
    from app.db.models import CaseMember

    make.db.add(CaseMember(case_id=case.id, user_id=viewer.user.id))
    make.db.commit()
    for path in ("reports", "export"):
        assert client.get(f"/v1/cases/{case.id}/{path}", headers=viewer.headers).status_code == 403
    assert (
        client.post(f"/v1/cases/{case.id}/reports", headers=viewer.headers, json={}).status_code
        == 403
    )

    _, findings_key = make.api_key(reviewer.tenant, "findings:read")
    _, cases_key = make.api_key(reviewer.tenant, "cases:read")
    assert (
        client.get(f"/v1/cases/{case.id}/export", headers=key_headers(findings_key)).status_code
        == 200
    )
    assert (
        client.get(f"/v1/cases/{case.id}/export", headers=key_headers(cases_key)).status_code == 403
    )
